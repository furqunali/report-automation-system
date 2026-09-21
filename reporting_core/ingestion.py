"""File ingestion layer.

Turns messy vendor CSV/XLSX exports into validated ``ReportRow`` records on
the canonical schema defined in :mod:`reporting_core.models`.

Real-world reports never share a header spelling: one store exports ``Qty``,
another ``Units Sold``, a third ``QTY ORDERED``. Rather than sprinkle string
comparisons through the codebase, this module drives a small, data-defined
*rules engine* (:class:`HeaderRules`) that maps any raw header onto a canonical
field. The readers themselves stay tiny and format-specific:

    * :func:`read_csv`  - pure standard-library ``csv`` (no third-party deps).
    * :func:`read_xlsx` - ``openpyxl``, imported lazily so CSV users never need
      it and a missing install fails with a clear, actionable message.

Both funnel through :func:`rows_from_records`, so CSV and XLSX inputs are
normalized identically. :func:`read_file` dispatches on the file extension.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator, Mapping, Sequence

from .models import ReportRow, Status

# Canonical fields, in the order ReportRow expects them.
CANONICAL_FIELDS: tuple[str, ...] = (
    "site",
    "category",
    "product",
    "quantity",
    "sales",
    "status",
)

# Raw-header alias (already normalized by :func:`normalize_key`) -> canonical
# field. Extend this as new vendor formats appear rather than editing readers.
DEFAULT_ALIASES: dict[str, str] = {
    # site / location
    "site": "site", "location": "site", "store": "site", "branch": "site",
    "store name": "site", "site name": "site",
    # category / grouping
    "category": "category", "group": "category", "product group": "category",
    "department": "category", "dept": "category", "class": "category",
    # product / item description
    "product": "product", "item": "product", "description": "product",
    "product name": "product", "product description": "product", "sku": "product",
    "item name": "product", "item description": "product",
    # quantity / movement
    "quantity": "quantity", "qty": "quantity", "units": "quantity",
    "qty sold": "quantity", "units sold": "quantity", "qty ordered": "quantity",
    "movement": "quantity", "count": "quantity",
    # sales / monetary amount
    "sales": "sales", "amount": "sales", "total": "sales", "cost": "sales",
    "net sales": "sales", "sales amount": "sales", "extended cost": "sales",
    "value": "sales", "revenue": "sales",
    # movement status
    "status": "status", "movement status": "status", "state": "status",
    "activity": "status",
}

# Raw status text (normalized) that means "no movement". Anything else that is
# present but unrecognized is treated as an active line.
_NO_MOVEMENT_TOKENS = frozenset({"no movement", "nomovement", "inactive", "dead", "stale"})


class IngestionError(Exception):
    """Base class for ingestion problems (bad file, unreadable content)."""


class UnsupportedFormatError(IngestionError):
    """Raised when a file's extension is not a recognized report format."""


class MissingDependencyError(IngestionError):
    """Raised when a format needs an optional dependency that is not installed."""


def normalize_key(raw: str | None) -> str:
    """Collapse a raw header to a comparable key.

    Lower-cased, punctuation/whitespace runs squeezed to single spaces, and
    trimmed - so ``"Sales ($)"``, ``"sales_$"`` and ``"  SALES  "`` all map to
    ``"sales"``.
    """
    return re.sub(r"[^a-z0-9]+", " ", (raw or "").strip().lower()).strip()


@dataclass
class HeaderRules:
    """Data-defined engine mapping raw headers to canonical fields."""

    aliases: dict[str, str] = field(default_factory=lambda: dict(DEFAULT_ALIASES))

    def add_rule(self, alias: str, field_name: str) -> None:
        """Register/override an alias. ``field_name`` must be canonical."""
        if field_name not in CANONICAL_FIELDS:
            raise ValueError(
                f"unknown canonical field {field_name!r}; "
                f"expected one of {CANONICAL_FIELDS}"
            )
        self.aliases[normalize_key(alias)] = field_name

    def resolve(self, raw_header: str | None) -> str | None:
        """Return the canonical field for a raw header, or ``None`` if unknown."""
        return self.aliases.get(normalize_key(raw_header))

    def map_headers(self, headers: Sequence[str | None]) -> dict[int, str]:
        """Map column *indices* to canonical fields.

        Later columns win on collision, mirroring how a human reading the sheet
        left-to-right would let a more specific column override an earlier one.
        """
        mapping: dict[int, str] = {}
        for index, header in enumerate(headers):
            canonical = self.resolve(header)
            if canonical is not None:
                mapping[index] = canonical
        return mapping


# Shared default engine so callers don't rebuild the alias table every read.
DEFAULT_RULES = HeaderRules()


def parse_number(value: object) -> float:
    """Parse a spreadsheet cell into a float.

    Handles thousands separators, currency symbols, blank cells and the
    accounting negative convention ``(1,234.50)`` -> ``-1234.50``. Values that
    cannot be parsed degrade to ``0.0`` rather than raising, because a single
    junk cell should not abort a multi-thousand-row ingest.
    """
    if isinstance(value, bool):  # bool is an int subclass; treat as unparseable
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("$", "").strip()
    if not text:
        return 0.0
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1].strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def normalize_status(value: object, *, quantity: float) -> Status:
    """Map a raw status cell onto the canonical ``Status`` literal.

    Explicit "no movement" style tokens win. When the status is blank we infer
    it from quantity: zero movement means ``"No Movement"``, otherwise
    ``"Active"``.
    """
    key = normalize_key(str(value)) if value is not None else ""
    if key in _NO_MOVEMENT_TOKENS or key.replace(" ", "") in _NO_MOVEMENT_TOKENS:
        return "No Movement"
    if not key:
        return "No Movement" if quantity <= 0 else "Active"
    return "Active"


def _record_to_row(record: Mapping[str, object]) -> ReportRow | None:
    """Build a ``ReportRow`` from a canonical-field record.

    Returns ``None`` for rows carrying neither a product nor a site, which are
    almost always spacer/subtotal lines rather than real data.
    """
    site = str(record.get("site", "") or "").strip()
    product = str(record.get("product", "") or "").strip()
    if not site and not product:
        return None

    quantity = parse_number(record.get("quantity", 0))
    sales = parse_number(record.get("sales", 0))
    return ReportRow(
        site=site,
        category=str(record.get("category", "") or "").strip(),
        product=product,
        quantity=quantity,
        sales=sales,
        status=normalize_status(record.get("status"), quantity=quantity),
    )


def rows_from_records(
    records: Iterable[Sequence[object]],
    headers: Sequence[str | None],
    rules: HeaderRules | None = None,
) -> list[ReportRow]:
    """Normalize positional rows (a header row + data rows) into ``ReportRow``s.

    This is the single funnel every reader shares: give it the header sequence
    plus an iterable of positional value rows and it applies the rules engine.
    """
    rules = rules or DEFAULT_RULES
    column_map = rules.map_headers(headers)
    rows: list[ReportRow] = []
    for values in records:
        record: dict[str, object] = {}
        for index, canonical in column_map.items():
            if index < len(values):
                record[canonical] = values[index]
        row = _record_to_row(record)
        if row is not None:
            rows.append(row)
    return rows


def read_csv(path: str | Path, rules: HeaderRules | None = None) -> list[ReportRow]:
    """Read a ``.csv`` report using only the standard library."""
    path = Path(path)
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        try:
            headers = next(reader)
        except StopIteration:
            return []
        return rows_from_records(reader, headers, rules)


def _iter_xlsx_records(path: Path) -> Iterator[tuple[list[str], Iterator[Sequence[object]]]]:
    """Yield ``(headers, data_rows)`` for each worksheet in an xlsx workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - exercised via monkeypatch
        raise MissingDependencyError(
            "Reading .xlsx/.xlsm files requires 'openpyxl'. "
            "Install it with:  pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            row_iter = worksheet.iter_rows(values_only=True)
            try:
                header_cells = next(row_iter)
            except StopIteration:
                continue  # empty sheet
            headers = [str(cell) if cell is not None else "" for cell in header_cells]
            yield headers, row_iter
    finally:
        workbook.close()


def read_xlsx(path: str | Path, rules: HeaderRules | None = None) -> list[ReportRow]:
    """Read a ``.xlsx``/``.xlsm`` report across all worksheets.

    Raises :class:`MissingDependencyError` with an actionable message if
    ``openpyxl`` is not installed.
    """
    path = Path(path)
    rows: list[ReportRow] = []
    for headers, data_rows in _iter_xlsx_records(path):
        rows.extend(rows_from_records(data_rows, headers, rules))
    return rows


def read_file(path: str | Path, rules: HeaderRules | None = None) -> list[ReportRow]:
    """Dispatch to the right reader based on the file extension."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path, rules)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path, rules)
    raise UnsupportedFormatError(
        f"unsupported report format {suffix!r} for {path.name!r}; "
        "expected .csv, .xlsx or .xlsm"
    )
